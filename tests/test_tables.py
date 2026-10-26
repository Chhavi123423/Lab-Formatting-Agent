"""
test_tables.py — Tests for table content in DOCX generation.
Verifies that tables are rendered as actual Word tables (not flat text),
column headers are present, and content is preserved exactly.
"""
import io
import pytest
from docx import Document as DocxDocument

from src.models import DocumentAST, HeadingNode, FormattingPlan
from src.formatting_planner import create_formatting_plan
from src.docx_generator import generate_docx_from_plan
from src.preview_renderer import _TablePlaceholder


def _make_ast_with_table():
    """Build a DocumentAST that contains a table placeholder node."""
    ast = DocumentAST(
        title="EXPERIMENT-1",
        subject_name="Data Structures",
        subject_code="CS-201",
        student_name="Test Student",
        roll_number="42",
        class_section="CS-A",
    )
    ast.nodes.append(HeadingNode(text="Observations", level=1, raw_text="Observations"))
    # Add a table placeholder directly
    tbl = _TablePlaceholder(
        headers=["Input", "Output", "Time (ms)"],
        rows=[["10", "True", "5"], ["20", "False", "12"]],
        caption="Table 1: Experimental Results",
    )
    ast.nodes.append(tbl)
    return ast


def test_table_node_rendered_as_word_table():
    """The generated docx must contain at least one Word table."""
    ast = _make_ast_with_table()
    plan = create_formatting_plan(ast)
    docx_bytes = generate_docx_from_plan(plan)

    doc = DocxDocument(io.BytesIO(docx_bytes))
    assert len(doc.tables) >= 1, "Expected at least one Word table in output"


def test_table_header_text_preserved():
    """Column headers must appear verbatim in the generated document."""
    ast = _make_ast_with_table()
    plan = create_formatting_plan(ast)
    docx_bytes = generate_docx_from_plan(plan)

    doc = DocxDocument(io.BytesIO(docx_bytes))
    all_table_text = ""
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    all_table_text += para.text + " "

    assert "Input" in all_table_text
    assert "Output" in all_table_text
    assert "Time (ms)" in all_table_text


def test_table_row_data_preserved():
    """Row data must be present in the Word table cells."""
    ast = _make_ast_with_table()
    plan = create_formatting_plan(ast)
    docx_bytes = generate_docx_from_plan(plan)

    doc = DocxDocument(io.BytesIO(docx_bytes))
    all_table_text = ""
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    all_table_text += para.text + " "

    assert "True" in all_table_text
    assert "False" in all_table_text


def test_empty_table_does_not_crash():
    """A table with no headers and no rows should not raise an exception."""
    ast = DocumentAST(title="Test")
    ast.nodes.append(HeadingNode(text="Data", level=1, raw_text="Data"))
    tbl = _TablePlaceholder(headers=[], rows=[], caption="")
    ast.nodes.append(tbl)
    plan = create_formatting_plan(ast)
    # Should complete without raising
    docx_bytes = generate_docx_from_plan(plan)
    assert len(docx_bytes) > 100

