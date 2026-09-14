import io
from docx import Document
from src.preprocessor import clean_input_text
from src.heading_detector import extract_sections_heuristically
from src.content_detector import build_document_ast
from src.profiles import get_university_practical_profile
from src.formatting_planner import create_formatting_plan
from src.docx_generator import generate_docx_from_plan
from src.validator import validate_generated_docx


def test_university_practical_file_generation():
    raw_sample = (
        "EXPERIMENT-3\n\n"
        "AIM: Write a C++ program to check whether a given string is grammar or not\n\n"
        "THEORY: A grammar in formal language theory is a set of production rules.\n\n"
        "CODE:\n"
        "#include <iostream>\n"
        "#include <string>\n"
        "using namespace std;\n"
        "int main() {\n"
        "    cout << \"Enter a string: \";\n"
        "    return 0;\n"
        "}\n\n"
        "OUTPUT:\n"
        "Enter a string: aaabbb"
    )

    cleaned = clean_input_text(raw_sample)
    sections = extract_sections_heuristically(cleaned)
    assert len(sections) >= 4

    ast = build_document_ast(
        raw_sections=sections,
        subject_name="CD PRACTICAL FILE",
        subject_code="CIC-351",
        student_name="Chhavi Choudhary",
        class_section="IT-2",
        roll_number="19",
        raw_input_text=cleaned
    )

    profile = get_university_practical_profile()
    plan = create_formatting_plan(ast, custom_profile=profile)
    docx_bytes = generate_docx_from_plan(plan)

    assert len(docx_bytes) > 2000

    # Inspect generated document properties
    doc = Document(io.BytesIO(docx_bytes))
    section = doc.sections[0]

    # Verify split header table
    header = section.header
    assert len(header.tables) > 0
    htable = header.tables[0]
    assert "CD PRACTICAL FILE" in htable.cell(0, 0).text
    assert "CIC-351" in htable.cell(0, 1).text

    # Verify student stamp in footer
    footer = section.footer
    footer_text = " ".join(p.text for p in footer.paragraphs)
    assert "Chhavi Choudhary" in footer_text
    assert "IT-2" in footer_text
    assert "19" in footer_text

    # Verify inline headings in body paragraphs
    body_texts = [p.text for p in doc.paragraphs]
    # Check that AIM: and its text are in the same paragraph
    aim_p = next((t for t in body_texts if "AIM:" in t), None)
    assert aim_p is not None
    assert "Write a C++ program" in aim_p

    # Validate output
    val_res = validate_generated_docx(docx_bytes, ast)
    assert val_res.status == "PASS"

