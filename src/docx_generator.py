import io
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from src.models import FormattingPlan


def apply_cell_styling(cell, bg_hex: str = "F4F4F4", border_hex: str = "D3D3D3"):
    """Applies background fill and clean borders to a table cell using Word XML."""
    tcPr = cell._tc.get_or_add_tcPr()

    # Background shading
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    tcPr.append(shd)

    # Subtle borders
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)


def set_monospaced_run_font(run, font_name: str, size_pt: float, color_rgb=(0, 0, 0)):
    """
    Forces all characters in a code run (including operators, brackets, and numbers)
    to strictly use the monospaced font without Word substituting body fonts.
    """
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(color_rgb[0], color_rgb[1], color_rgb[2])

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)


def add_page_number_to_run(run):
    """Inserts a dynamic Word PAGE field into a run."""
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    run._r.append(fldSimple)


def generate_docx_from_plan(plan: FormattingPlan) -> bytes:
    """
    Deterministically renders a Word .docx document from a FormattingPlan.
    Supports preset profiles and dynamic AI-extracted profiles.
    """
    doc = Document()
    section = doc.sections[0]

    # 1. Page Configuration
    p_cfg = plan.page_config
    section.page_width = Inches(p_cfg.get("width_inches", 8.27))
    section.page_height = Inches(p_cfg.get("height_inches", 11.69))
    section.top_margin = Inches(p_cfg.get("margin_top_inches", 1.0))
    section.bottom_margin = Inches(p_cfg.get("margin_bottom_inches", 1.0))
    section.left_margin = Inches(p_cfg.get("margin_left_inches", 1.0))
    section.right_margin = Inches(p_cfg.get("margin_right_inches", 1.0))

    # 2. Header
    header = section.header
    header_layout = plan.header_config.get("layout", "")
    h_font = plan.header_config.get("font", "Times New Roman")
    h_size = plan.header_config.get("size_pt", 11)
    h_color = plan.header_config.get("color_rgb", (0, 0, 0))

    if header_layout == "TWO_COLUMN_SPLIT" or (plan.subject_name and plan.subject_code):
        # Create a 1-row, 2-column borderless table in the header
        htable = header.add_table(rows=1, cols=2, width=Inches(6.27))
        htable.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Left cell: Subject Name (e.g. CD PRACTICAL FILE)
        cell_l = htable.cell(0, 0)
        cell_l.width = Inches(3.2)
        p_l = cell_l.paragraphs[0]
        p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_l = p_l.add_run(plan.subject_name or "PRACTICAL FILE")
        run_l.font.name = h_font
        run_l.font.size = Pt(h_size)
        run_l.font.color.rgb = RGBColor(h_color[0], h_color[1], h_color[2])

        # Right cell: Subject Code (e.g. CIC-351)
        cell_r = htable.cell(0, 1)
        cell_r.width = Inches(3.07)
        p_r = cell_r.paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_r = p_r.add_run(plan.subject_code or "")
        run_r.font.name = h_font
        run_r.font.size = Pt(h_size)
        run_r.font.color.rgb = RGBColor(h_color[0], h_color[1], h_color[2])
    elif header_layout == "CENTERED" and plan.subject_name:
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hrun = hp.add_run(plan.subject_name)
        hrun.font.name = h_font
        hrun.font.size = Pt(h_size)
        hrun.font.color.rgb = RGBColor(h_color[0], h_color[1], h_color[2])
    elif plan.student_name or plan.roll_number:
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        meta_runs = []
        if plan.student_name:
            meta_runs.append(f"Name: {plan.student_name}")
        if plan.roll_number:
            meta_runs.append(f"Roll No: {plan.roll_number}")
        
        hrun = hp.add_run(" | ".join(meta_runs))
        hrun.font.name = plan.typography.get("header", {}).get("font", "Calibri")
        hrun.font.size = Pt(plan.typography.get("header", {}).get("size_pt", 9))
        hrun.font.color.rgb = RGBColor(128, 128, 128)

    # 3. Footer Handling (Flexible, Not Hardcoded)
    footer = section.footer
    footer_layout = plan.footer_config.get("layout", "PAGE_NUMBER")
    f_font = plan.footer_config.get("font", "Times New Roman")
    f_size = plan.footer_config.get("size_pt", 10)
    f_color = plan.footer_config.get("color_rgb", (100, 100, 100))

    if footer_layout == "STUDENT_STAMP" and (plan.student_name or plan.roll_number):
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.paragraph_format.line_spacing = 1.0
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(2)

        stamp_lines = []
        if plan.student_name:
            stamp_lines.append(plan.student_name)
        if plan.class_section:
            stamp_lines.append(plan.class_section)
        if plan.roll_number:
            stamp_lines.append(f"Roll no: - {plan.roll_number}")

        for s_idx, s_line in enumerate(stamp_lines):
            curr_p = fp if s_idx == 0 else footer.add_paragraph()
            curr_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            curr_p.paragraph_format.line_spacing = 1.0
            curr_p.paragraph_format.space_before = Pt(0)
            curr_p.paragraph_format.space_after = Pt(2)
            s_run = curr_p.add_run(s_line)
            s_run.font.name = f_font
            s_run.font.size = Pt(f_size)
            s_run.font.color.rgb = RGBColor(0, 0, 0)

    elif footer_layout == "PAGE_NUMBER":
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Page ")
        frun.font.name = f_font
        frun.font.size = Pt(f_size)
        frun.font.color.rgb = RGBColor(f_color[0], f_color[1], f_color[2])
        add_page_number_to_run(frun)

    elif footer_layout == "CUSTOM_TEXT":
        custom_text = plan.footer_config.get("text") or "Lab Record"
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run(custom_text)
        frun.font.name = f_font
        frun.font.size = Pt(f_size)
        frun.font.italic = True
        frun.font.color.rgb = RGBColor(f_color[0], f_color[1], f_color[2])

    elif footer_layout == "NONE":
        # Completely empty footer
        pass

    else:
        # Default: centered page number
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Page ")
        frun.font.name = f_font
        frun.font.size = Pt(f_size)
        frun.font.color.rgb = RGBColor(f_color[0], f_color[1], f_color[2])
        add_page_number_to_run(frun)

    # 4. Render Nodes
    for item in plan.nodes:
        node_type = item["type"]
        style = item.get("style", {})

        if node_type == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(style.get("space_before_pt", 6))
            p.paragraph_format.space_after = Pt(style.get("space_after_pt", 12))
            title_text = item["text"].strip()
            if style.get("uppercase", True):
                title_text = title_text.upper()
            run = p.add_run(title_text)
            run.font.name = style.get("font", "Times New Roman")
            run.font.size = Pt(style.get("size_pt", 14))
            run.font.bold = True
            c = style.get("color_rgb", (0, 0, 0))
            run.font.color.rgb = RGBColor(c[0], c[1], c[2])

        elif node_type == "heading":
            is_inline = item.get("is_inline", False)
            inline_text = item.get("inline_text", "")

            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = style.get("keep_with_next", True)
            p.paragraph_format.space_before = Pt(style.get("space_before_pt", 8))
            p.paragraph_format.space_after = Pt(style.get("space_after_pt", 4))
            
            # Heading text (e.g. "AIM:" or "THEORY:")
            head_label = item["text"].strip().rstrip(":")
            if style.get("uppercase", True):
                head_label = head_label.upper()
            head_str = head_label + ": "
            run_head = p.add_run(head_str)
            run_head.font.name = style.get("font", "Times New Roman")
            run_head.font.size = Pt(style.get("size_pt", 12))
            run_head.font.bold = True
            run_head.font.italic = style.get("italic", False)
            c = style.get("color_rgb", (0, 0, 0))
            run_head.font.color.rgb = RGBColor(c[0], c[1], c[2])

            # If inline, append the content immediately on the same paragraph
            if is_inline and inline_text:
                run_body = p.add_run(inline_text)
                b_style = item.get("body_style", {})
                run_body.font.name = b_style.get("font", style.get("font", "Times New Roman"))
                run_body.font.size = Pt(b_style.get("size_pt", 12))
                run_body.font.bold = False
                p.paragraph_format.line_spacing = b_style.get("line_spacing", 1.15)
                p.paragraph_format.space_after = Pt(b_style.get("space_after_pt", 6))

        elif node_type == "paragraph":
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = style.get("line_spacing", 1.15)
            p.paragraph_format.space_after = Pt(style.get("space_after_pt", 6))
            
            run = p.add_run(item["text"])
            run.font.name = style.get("font", "Times New Roman")
            run.font.size = Pt(style.get("size_pt", 12))
            c = style.get("color_rgb", (0, 0, 0))
            run.font.color.rgb = RGBColor(c[0], c[1], c[2])

        elif node_type == "code_block":
            boxed = style.get("boxed", True)
            code_lines = item["code"].split("\n")
            code_font = style.get("font", "Consolas")
            code_size = style.get("size_pt", 10)

            if boxed:
                # Render code inside a single-cell table with background shading
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False

                cell = table.cell(0, 0)
                cell.width = Inches(6.2)
                apply_cell_styling(cell, bg_hex=style.get("bg_color_hex", "F4F4F4"), border_hex=style.get("border_color_hex", "D3D3D3"))

                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(4)
                cp.paragraph_format.space_after = Pt(4)
                cp.paragraph_format.line_spacing = 1.0

                for idx, cline in enumerate(code_lines):
                    if idx > 0:
                        cp = cell.add_paragraph()
                        cp.paragraph_format.space_before = Pt(0)
                        cp.paragraph_format.space_after = Pt(0)
                        cp.paragraph_format.line_spacing = 1.0

                    crun = cp.add_run(cline)
                    set_monospaced_run_font(crun, code_font, 9.5, (30, 30, 30))

                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(6)
            else:
                # Clean unboxed monospaced code as in University Practical format
                for idx, cline in enumerate(code_lines):
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.2)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(1 if idx < len(code_lines) - 1 else 6)
                    crun = p.add_run(cline)
                    set_monospaced_run_font(crun, code_font, code_size, (0, 0, 0))

        elif node_type == "list":
            items = item.get("items", [])
            is_ordered = item.get("is_ordered", False)
            for idx, lit in enumerate(items, 1):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Inches(0.3)
                
                bullet_prefix = f"{idx}. " if is_ordered else "•  "
                brun = p.add_run(bullet_prefix)
                brun.font.bold = True
                brun.font.name = style.get("font", "Times New Roman")
                brun.font.size = Pt(style.get("size_pt", 12))

                trun = p.add_run(lit)
                trun.font.name = style.get("font", "Times New Roman")
                trun.font.size = Pt(style.get("size_pt", 12))

        elif node_type == "image":
            image_bytes = item.get("image_bytes")
            if image_bytes:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(style.get("space_before_pt", 6))
                p.paragraph_format.space_after = Pt(2)
                
                max_w = style.get("max_width_inches", 5.5)
                image_stream = io.BytesIO(image_bytes)
                p.add_run().add_picture(image_stream, width=Inches(max_w))

                caption = item.get("caption")
                if caption:
                    cap_p = doc.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_after = Pt(style.get("space_after_pt", 6))
                    crun = cap_p.add_run(caption)
                    crun.font.name = style.get("caption_font", "Times New Roman")
                    crun.font.size = Pt(style.get("caption_size_pt", 9.5))
                    crun.font.italic = True
                    crun.font.color.rgb = RGBColor(100, 100, 100)

    # Return DOCX file content as bytes
    out_stream = io.BytesIO()
    doc.save(out_stream)
    return out_stream.getvalue()
